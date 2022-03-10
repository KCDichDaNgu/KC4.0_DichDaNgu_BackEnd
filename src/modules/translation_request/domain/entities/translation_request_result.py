import io
import openpyxl
from pptx.api import Presentation
from pydantic import Field
from typing import Any, Optional, get_args, IO
from docx import Document

import pickle
import os

from sanic.request import File

from core.base_classes.entity import Entity
from core.value_objects import ID

from modules.task.domain.entities.task_result import TaskResultEntity, TaskResultProps

from infrastructure.configs.translation_task import (
    FILE_TRANSLATION_FOLDER_PATH,
    FILE_TRANSLATION_TASKS,
    AllowedFileTranslationExtensionEnum,
    get_file_translation_binary_progress_file_name, 
    get_file_translation_source_file_name,
    get_file_translation_file_path
)

from infrastructure.configs.message import MESSAGES
from infrastructure.configs.main import StatusCodeEnum

from interface_adapters.dtos.base_response import BaseResponse

from core.utils.file import extract_file_extension, get_full_path

class TranslationRequestResultProps(TaskResultProps):
    
    task_id: ID = Field(...)
    step: str = Field(...)
    file_path: Optional[str]

class TranslationRequestResultEntity(
    TaskResultEntity, 
    Entity[TranslationRequestResultProps]
):

    @property
    def props_klass(self):
        return get_args(self.__orig_bases__[1])[0]

    def __create_file_translation_task_folder(self):

        real_file_translation_task_folder = get_full_path(f'{FILE_TRANSLATION_FOLDER_PATH}/{self.props.task_id.value}')

        if not os.path.isdir(real_file_translation_task_folder):

            os.makedirs(real_file_translation_task_folder)

        return real_file_translation_task_folder

    async def create_required_files_for_file_translation_task(
        self,
        binary: IO,
        original_file_ext: AllowedFileTranslationExtensionEnum
    ):    
        # if self.props.task_name not in FILE_TRANSLATION_TASKS:
        #     raise ValueError('Translation task is not file translation')
        original_file_name = f'{get_file_translation_source_file_name()}.{original_file_ext}'
        original_file_path = get_file_translation_file_path(self.props.task_id.value, original_file_name)
        original_file_full_path = get_full_path(original_file_path)

        binary_progress_file_name = f'{get_file_translation_binary_progress_file_name()}'
        binary_progress_file_path = get_file_translation_file_path(self.props.task_id.value, binary_progress_file_name)
        binary_progress_file_full_path = get_full_path(binary_progress_file_path)

        self.__create_file_translation_task_folder()

        try:
            if original_file_ext == 'docx':          
                file = Document(binary)
                
            elif original_file_ext == 'pptx':            
                file = Presentation(binary)
                
            elif original_file_ext == 'xlsx':            
                file = openpyxl.load_workbook(binary)

            file.save(original_file_full_path)

            with open(binary_progress_file_full_path, 'wb') as outp:
                pickle.dump(binary, outp, pickle.HIGHEST_PROTOCOL)

        except Exception as e:
            return BaseResponse(
                code=StatusCodeEnum.failed.value,
                data=dict(
                    original_file_full_path=None,
                    binary_progress_file_full_path=None
                ),
                message=MESSAGES['failed']
            )
        
        return BaseResponse(
            code=StatusCodeEnum.success.value,
            data=dict(
                original_file_full_path=original_file_full_path,
                binary_progress_file_full_path=binary_progress_file_full_path
            ),
            message=MESSAGES['success']
        )

    async def create_required_files_for_txt_file_translation_task(self, original_file: File):
        self.__create_file_translation_task_folder()
        
        original_file_ext = extract_file_extension(original_file.name)
        original_file_name = f'{get_file_translation_source_file_name()}.{original_file_ext}'
        original_file_path = get_file_translation_file_path(self.props.task_id.value, original_file_name)
        original_file_full_path = get_full_path(original_file_path)
        try:
            with open(original_file_full_path, 'wb+') as outp:

                outp.write(original_file.body)

                outp.close()
        
        except Exception as e:
            return BaseResponse(
                code=StatusCodeEnum.failed.value,
                data=dict(
                    original_file_full_path=None
                ),
                message=MESSAGES['failed']
            )
        
        return BaseResponse(
                code=StatusCodeEnum.success.value,
                data=dict(
                    original_file_full_path=original_file_full_path
                ),
                message=MESSAGES['success']
            )